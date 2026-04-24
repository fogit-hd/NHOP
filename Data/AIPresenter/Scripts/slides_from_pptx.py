"""Build Temp/slides_iterate.tsv for UiPath from available slide sources.

Priority:
1) Configured .pptx path in Data/AIPresenter/Config/ai_settings.json
2) Ordered .pptx list in Data/AIPresenter/Input/pptx_run_order.xlsx
    or Data/AIPresenter/Input/pptx_run_order_template.xlsx
3) First .pptx in Data/AIPresenter/Input/pptx_file_here
4) Data/AIPresenter/Input/slides.xlsx (legacy fallback)

TSV columns (tab-separated): index, title, image_path, relative_path
Rows are sorted by slide_index ascending.
"""
from __future__ import annotations

import argparse
import importlib
import json
import os
import shutil
import sys
from typing import Any

RUN_ORDER_DEFAULT = os.path.join("Data", "AIPresenter", "Input", "pptx_run_order.xlsx")
RUN_ORDER_TEMPLATE_XLSX_DEFAULT = os.path.join(
    "Data", "AIPresenter", "Input", "pptx_run_order_template.xlsx"
)
SUPPORTED_WORKBOOK_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
PPTX_SOURCE_ERROR_FILE = os.path.join(
    "Data", "AIPresenter", "Temp", "pptx_source_resolution_error.txt"
)
PPTX_EXPORT_DIAGNOSTICS_FILE = os.path.join(
    "Data", "AIPresenter", "Temp", "pptx_export_diagnostics.txt"
)
DEFAULT_EXPORT_WIDTH = 2560
DEFAULT_EXPORT_HEIGHT = 1440


def configure_stdio_utf8_fallback() -> None:
    for stream_name in ("stdout", "stderr"):
        stream = getattr(sys, stream_name, None)
        if stream is None:
            continue
        reconfigure = getattr(stream, "reconfigure", None)
        if not callable(reconfigure):
            continue
        try:
            reconfigure(encoding="utf-8", errors="backslashreplace")
        except Exception:
            pass


configure_stdio_utf8_fallback()


def clear_pptx_export_root(root: str) -> str:
    out_dir = os.path.join(root, "Data", "AIPresenter", "Temp", "pptx_slides")
    os.makedirs(out_dir, exist_ok=True)

    for name in os.listdir(out_dir):
        target = os.path.join(out_dir, name)
        try:
            if os.path.isdir(target):
                shutil.rmtree(target)
            else:
                os.remove(target)
        except OSError as e:
            print(f"WARN Could not remove old export item: {target} ({e})", file=sys.stderr)
    return out_dir


def sanitize_name(name: str) -> str:
    cleaned = "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in name)
    cleaned = cleaned.strip("_")
    return cleaned or "presentation"


def sanitize_tsv_field(value: Any) -> str:
    """Flatten tabs/newlines so each TSV row always remains 4 physical columns."""
    text = str(value or "")
    return text.replace("\r\n", " ").replace("\n", " ").replace("\r", " ").replace("\t", " ").strip()


def parse_export_dimension(value: Any, default_value: int) -> int:
    try:
        parsed = int(float(str(value).strip()))
        if parsed >= 640:
            return parsed
    except (TypeError, ValueError):
        pass
    return default_value


def find_first_pptx(folder: str) -> str:
    if not os.path.isdir(folder):
        return ""
    for name in sorted(os.listdir(folder)):
        low = name.lower()
        if low.endswith(".pptx") and not low.startswith("~$"):
            candidate = os.path.normpath(os.path.join(folder, name))
            if os.path.isfile(candidate):
                return candidate
    return ""


def list_pptx_files(folder: str) -> list[str]:
    files: list[str] = []
    if not os.path.isdir(folder):
        return files
    for name in sorted(os.listdir(folder)):
        low = name.lower()
        if low.endswith(".pptx") and not low.startswith("~$"):
            candidate = os.path.normpath(os.path.join(folder, name))
            if os.path.isfile(candidate):
                files.append(candidate)
    return files


def format_path_for_message(root: str, path: str) -> str:
    try:
        rel = os.path.relpath(path, root)
        return rel.replace("\\", "/")
    except ValueError:
        return path.replace("\\", "/")


def project_relative_path_or_original(root: str, path_value: str) -> str:
    raw = os.path.normpath(str(path_value or "").strip())
    if not raw:
        return ""
    if not os.path.isabs(raw):
        return raw
    try:
        rel = os.path.relpath(raw, root)
    except ValueError:
        return raw
    rel_norm = os.path.normpath(rel)
    if rel_norm.startswith("..") or os.path.isabs(rel_norm):
        return raw
    return rel_norm


def build_run_order_mismatch_message(
    root: str,
    queue_path: str,
    unresolved_entries: list[str],
    folder_pptx_files: list[str],
) -> str:
    unresolved_unique: list[str] = []
    seen: set[str] = set()
    for raw in unresolved_entries:
        key = raw.strip().lower()
        if not key or key in seen:
            continue
        seen.add(key)
        unresolved_unique.append(raw.strip())

    lines: list[str] = [
        "PPTX source mismatch: run-order file and PPTX folder are not aligned.",
        f"Run-order file: {format_path_for_message(root, queue_path)}",
        "",
        "Entries listed in run-order but not found on disk:",
    ]

    if unresolved_unique:
        for raw in unresolved_unique[:30]:
            lines.append(f"- {raw}")
        if len(unresolved_unique) > 30:
            lines.append(f"- ... and {len(unresolved_unique) - 30} more")
    else:
        lines.append("- (none)")

    lines.extend(["", "PPTX files currently found in Data/AIPresenter/Input/pptx_file_here:"])
    if folder_pptx_files:
        for path in folder_pptx_files[:30]:
            lines.append(f"- {format_path_for_message(root, path)}")
        if len(folder_pptx_files) > 30:
            lines.append(f"- ... and {len(folder_pptx_files) - 30} more")
    else:
        lines.append("- (none)")

    lines.extend(
        [
            "",
            "Fix:",
            "- Update pptx_run_order_template.xlsx (or pptx_run_order.xlsx).",
            "- Ensure each pptx_path exactly matches an existing .pptx file name/path.",
        ]
    )
    return "\n".join(lines)


def build_run_order_no_enabled_message(root: str, queue_path: str) -> str:
    lines: list[str] = [
        "No PPTX file is enabled in the run-order file.",
        f"Run-order file: {format_path_for_message(root, queue_path)}",
        "",
        "Fix:",
        "- Set enabled to TRUE for at least one pptx_path row.",
        "- Ensure each enabled row points to an existing .pptx file.",
    ]
    return "\n".join(lines)


def clear_pptx_source_error(root: str) -> None:
    path = os.path.join(root, PPTX_SOURCE_ERROR_FILE)
    try:
        if os.path.isfile(path):
            os.remove(path)
    except OSError:
        pass


def write_pptx_source_error(root: str, message: str) -> None:
    path = os.path.join(root, PPTX_SOURCE_ERROR_FILE)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write((message or "").strip() + "\n")


def clear_pptx_export_diagnostics(root: str) -> None:
    path = os.path.join(root, PPTX_EXPORT_DIAGNOSTICS_FILE)
    try:
        if os.path.isfile(path):
            os.remove(path)
    except OSError:
        pass


def append_pptx_export_diagnostics(
    root: str,
    pptx_path: str,
    out_dir: str,
    slide_count: int,
    export_width: int,
    export_height: int,
    notes: list[str],
) -> None:
    path = os.path.join(root, PPTX_EXPORT_DIAGNOSTICS_FILE)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    has_existing = os.path.isfile(path) and os.path.getsize(path) > 0
    lines = [
        f"pptx={format_path_for_message(root, pptx_path)}",
        f"output_dir={format_path_for_message(root, out_dir)}",
        f"requested_slide_count={slide_count}",
        f"export_resolution={export_width}x{export_height}",
        f"python_executable={sys.executable}",
        f"python_version={sys.version.split()[0]}",
    ]
    lines.extend(str(note) for note in notes if str(note).strip())
    with open(path, "a", encoding="utf-8") as f:
        if has_existing:
            f.write("\n" + ("=" * 80) + "\n")
        f.write("\n".join(lines).strip() + "\n")


def format_exception_message(error: BaseException) -> str:
    return f"{type(error).__name__}: {error}"


def export_with_powerpoint_backend(
    client_module: Any,
    backend_name: str,
    pptx_path: str,
    out_dir: str,
    slide_count: int,
    export_width: int,
    export_height: int,
) -> tuple[list[str], str]:
    app = None
    pres = None
    exported: list[str] = [""] * max(0, slide_count)
    try:
        create_object = getattr(client_module, "CreateObject", None)
        if callable(create_object):
            app = create_object("PowerPoint.Application")
        else:
            dispatch = getattr(client_module, "Dispatch", None)
            if callable(dispatch):
                app = dispatch("PowerPoint.Application")
            else:
                raise RuntimeError(
                    f"{backend_name} module does not expose CreateObject/Dispatch for PowerPoint."
                )
        app.Visible = 1
        pres = app.Presentations.Open(pptx_path, False, False, False)
        total = int(pres.Slides.Count)
        if total > len(exported):
            exported.extend([""] * (total - len(exported)))
        for idx in range(1, total + 1):
            out_png = os.path.join(out_dir, f"slide_{idx:03d}.png")
            slide = pres.Slides.Item(idx)
            slide.Export(out_png, "PNG", export_width, export_height)
            if os.path.isfile(out_png):
                exported[idx - 1] = os.path.normpath(out_png)
        return exported, ""
    except Exception as error:
        return exported, f"{backend_name} export failed: {format_exception_message(error)}"
    finally:
        try:
            if pres is not None:
                pres.Close()
        except Exception:
            pass
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass


def parse_enabled(value: Any) -> bool:
    if value is None:
        return True
    text = str(value).strip().lower()
    if text == "":
        return True
    return text not in {"0", "false", "no", "n", "off", "disabled"}


def resolve_pptx_candidate(root: str, input_dir: str, raw_path: str) -> str:
    value = (raw_path or "").strip().strip('"')
    if not value:
        return ""

    candidates: list[str] = []
    if os.path.isabs(value):
        candidates.append(value)
    else:
        candidates.append(os.path.join(root, value))
        candidates.append(os.path.join(input_dir, "pptx_file_here", value))

    for candidate in candidates:
        normalized = os.path.normpath(candidate)
        base_name = os.path.basename(normalized).lower()
        if base_name.startswith("~$") or not base_name.endswith(".pptx"):
            continue
        if os.path.isfile(normalized):
            return normalized
    return ""


def read_pptx_run_order(root: str, input_dir: str, queue_path: str) -> tuple[list[str], int, list[str]]:
    if not os.path.isfile(queue_path):
        return [], 0, []

    try:
        from openpyxl import load_workbook
    except Exception:
        print(
            "WARN pptx_run_order.xlsx found but openpyxl is missing. "
            "Install dependencies to use ordered multi-PPTX mode.",
            file=sys.stderr,
        )
        return [], 0, []

    try:
        workbook = load_workbook(queue_path, data_only=True, read_only=True)
        sheet = workbook.active
    except Exception as e:
        print(f"WARN Could not read PPTX run-order file: {queue_path} ({e})", file=sys.stderr)
        return [], 0, []

    header_row = None
    for values in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
        header_row = list(values)
        break

    if not header_row:
        return [], 0, []

    def norm_header(value: Any) -> str:
        return str(value or "").strip().lower().replace(" ", "_")

    headers = [norm_header(v) for v in header_row]
    path_aliases = {"pptx_path", "file_path", "path", "pptx", "file", "filename", "pptx_file"}
    order_aliases = {"run_order", "order", "stt", "index", "priority"}
    enabled_aliases = {"enabled", "is_enabled", "active", "use"}

    path_idx = next((i for i, h in enumerate(headers) if h in path_aliases), -1)
    order_idx = next((i for i, h in enumerate(headers) if h in order_aliases), -1)
    enabled_idx = next((i for i, h in enumerate(headers) if h in enabled_aliases), -1)

    if path_idx < 0:
        print(
            "WARN pptx_run_order.xlsx does not contain a path column. "
            "Expected one of: pptx_path, file_path, path, pptx.",
            file=sys.stderr,
        )
        return [], 0, []

    items: list[tuple[int, str]] = []
    unresolved_paths: list[str] = []
    rows_with_path = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = list(row)

        raw_path = ""
        if path_idx < len(values) and values[path_idx] is not None:
            raw_path = str(values[path_idx]).strip()
        if not raw_path:
            continue

        rows_with_path += 1

        enabled = True
        if enabled_idx >= 0 and enabled_idx < len(values):
            enabled = parse_enabled(values[enabled_idx])
        if not enabled:
            continue

        order_value = rows_with_path
        if order_idx >= 0 and order_idx < len(values) and values[order_idx] not in (None, ""):
            try:
                order_value = int(float(str(values[order_idx]).strip()))
            except ValueError:
                order_value = rows_with_path

        resolved = resolve_pptx_candidate(root, input_dir, raw_path)
        if not resolved:
            print(f"WARN queue path not found: {raw_path}", file=sys.stderr)
            unresolved_paths.append(raw_path)
            continue

        items.append((order_value, resolved))

    items.sort(key=lambda it: (it[0], it[1].lower()))
    ordered_paths: list[str] = []
    seen: set[str] = set()
    for _, path in items:
        key = path.lower()
        if key in seen:
            continue
        seen.add(key)
        ordered_paths.append(path)

    return ordered_paths, rows_with_path, unresolved_paths


def resolve_pptx_paths(root: str, cfg: dict) -> list[str]:
    input_dir = os.path.join(root, "Data", "AIPresenter", "Input")
    preferred_dir = os.path.join(input_dir, "pptx_file_here")

    configured = (
        (cfg.get("presentation_pptx_path") or "").strip()
        or (cfg.get("pptx_input_path") or "").strip()
        or (cfg.get("slide_pptx_path") or "").strip()
    )
    if configured:
        candidate = resolve_pptx_candidate(root, input_dir, configured)
        return [candidate] if candidate else []

    queue_setting = (cfg.get("pptx_run_order_file") or "").strip()
    queue_candidates: list[str] = []
    if queue_setting:
        queue_candidates.append(
            os.path.normpath(queue_setting)
            if os.path.isabs(queue_setting)
            else os.path.normpath(os.path.join(root, queue_setting))
        )
    else:
        queue_candidates.append(os.path.join(root, RUN_ORDER_DEFAULT))
        queue_candidates.append(os.path.join(root, RUN_ORDER_TEMPLATE_XLSX_DEFAULT))

    queue_rows_detected = 0
    queue_path_with_rows = ""
    unresolved_queue_entries: list[str] = []
    for queue_path in queue_candidates:
        ext = os.path.splitext(queue_path)[1].lower()
        if ext == ".csv":
            raise RuntimeError(
                "CSV run-order is no longer supported. "
                "Use Excel workbook .xlsx for pptx_run_order_file."
            )
        if ext and ext not in SUPPORTED_WORKBOOK_EXTENSIONS:
            print(
                f"WARN Unsupported run-order extension '{ext}' for file: {queue_path}. "
                "Expected an Excel workbook (.xlsx/.xlsm/.xltx/.xltm).",
                file=sys.stderr,
            )
            continue
        queue_paths, queue_rows, unresolved = read_pptx_run_order(root, input_dir, queue_path)

        if queue_rows > 0 and unresolved:
            preferred_files = list_pptx_files(preferred_dir)
            raise RuntimeError(
                build_run_order_mismatch_message(
                    root,
                    queue_path,
                    unresolved,
                    preferred_files,
                )
            )

        if queue_rows > 0 and not queue_paths and not unresolved:
            raise RuntimeError(build_run_order_no_enabled_message(root, queue_path))

        if queue_paths:
            return queue_paths
        if queue_rows > 0 and not queue_path_with_rows:
            queue_path_with_rows = queue_path
        queue_rows_detected = max(queue_rows_detected, queue_rows)
        unresolved_queue_entries.extend(unresolved)

    if queue_rows_detected > 0:
        preferred_files = list_pptx_files(preferred_dir)

        if unresolved_queue_entries:
            queue_path_for_message = queue_path_with_rows or queue_candidates[0]
            raise RuntimeError(
                build_run_order_mismatch_message(
                    root,
                    queue_path_for_message,
                    unresolved_queue_entries,
                    preferred_files,
                )
            )

        print(
            "WARN PPTX run-order file exists but no valid entries were resolved. "
            "Falling back to folder auto-detection.",
            file=sys.stderr,
        )

    preferred_dir = os.path.join(input_dir, "pptx_file_here")
    preferred_candidate = find_first_pptx(preferred_dir)
    if preferred_candidate:
        return [preferred_candidate]

    return []


def load_rows_from_excel(root: str, workbook_path: str) -> list[tuple[int, str, str, str]]:
    rows: list[tuple[int, str, str, str]] = []
    if not os.path.isfile(workbook_path):
        return rows

    try:
        from openpyxl import load_workbook
    except Exception:
        print(
            "WARN slides.xlsx found but openpyxl is missing. "
            "Install dependencies to use Excel slide source mode.",
            file=sys.stderr,
        )
        return rows

    workbook = None
    try:
        workbook = load_workbook(workbook_path, data_only=True, read_only=True)
        sheet = workbook.active
        header_row = None
        for values in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
            header_row = list(values)
            break
        if not header_row:
            return rows

        def norm_header(value: Any) -> str:
            return str(value or "").strip().lower().replace(" ", "_")

        headers = [norm_header(v) for v in header_row]
        idx_aliases = {"slide_index", "index", "slide", "slide_no", "slide_number"}
        title_aliases = {"slide_title", "title", "slide_name", "name"}
        path_aliases = {"slide_path", "path", "image_path", "file_path"}

        idx_col = next((i for i, h in enumerate(headers) if h in idx_aliases), -1)
        title_col = next((i for i, h in enumerate(headers) if h in title_aliases), -1)
        path_col = next((i for i, h in enumerate(headers) if h in path_aliases), -1)
        if idx_col < 0 or path_col < 0:
            print(
                "WARN slides.xlsx missing required columns. "
                "Expected at least: slide_index and slide_path.",
                file=sys.stderr,
            )
            return rows

        for values in sheet.iter_rows(min_row=2, values_only=True):
            row = list(values)
            idx_raw = ""
            rel_raw = ""
            title_raw = ""
            if idx_col < len(row) and row[idx_col] is not None:
                idx_raw = str(row[idx_col]).strip()
            if path_col < len(row) and row[path_col] is not None:
                rel_raw = str(row[path_col]).strip().strip('"')
            if title_col >= 0 and title_col < len(row) and row[title_col] is not None:
                title_raw = str(row[title_col]).strip()

            if not idx_raw or not rel_raw:
                continue
            try:
                idx = int(float(idx_raw))
            except ValueError:
                continue
            stored_path = project_relative_path_or_original(root, rel_raw)
            rows.append((idx, title_raw, stored_path, stored_path))
    except Exception as e:
        print(f"WARN Could not read slides workbook: {workbook_path} ({e})", file=sys.stderr)
    finally:
        try:
            if workbook is not None:
                workbook.close()
        except Exception:
            pass

    rows.sort(key=lambda x: x[0])
    return rows


def export_pptx_slides_to_images(
    root: str,
    pptx_path: str,
    slide_count: int,
    deck_order: int,
    export_width: int,
    export_height: int,
) -> list[str]:
    out_root = os.path.join(root, "Data", "AIPresenter", "Temp", "pptx_slides")
    deck_name = sanitize_name(os.path.splitext(os.path.basename(pptx_path))[0])
    out_dir = os.path.join(out_root, f"deck_{deck_order:03d}_{deck_name}")
    os.makedirs(out_dir, exist_ok=True)

    notes: list[str] = []
    exported: list[str] = [""] * max(0, slide_count)
    backend_used = ""

    probe_path = os.path.join(out_dir, ".__write_probe__.tmp")
    try:
        with open(probe_path, "w", encoding="utf-8") as probe:
            probe.write("ok")
        os.remove(probe_path)
        notes.append("output_dir_write_test=OK")
    except Exception as error:
        notes.append(f"output_dir_write_test=FAIL ({format_exception_message(error)})")
        try:
            if os.path.isfile(probe_path):
                os.remove(probe_path)
        except OSError:
            pass

    try:
        with open(pptx_path, "rb") as test_in:
            test_in.read(1)
        notes.append("pptx_read_test=OK")
    except Exception as error:
        notes.append(f"pptx_read_test=FAIL ({format_exception_message(error)})")

    comtypes_client = None
    try:
        comtypes_client = importlib.import_module("comtypes.client")
        notes.append("backend_comtypes_import=OK")
    except Exception as error:
        notes.append(f"backend_comtypes_import=FAIL ({format_exception_message(error)})")

    if comtypes_client is not None:
        exported, backend_error = export_with_powerpoint_backend(
            comtypes_client,
            "comtypes",
            pptx_path,
            out_dir,
            slide_count,
            export_width,
            export_height,
        )
        if backend_error:
            notes.append(f"backend_comtypes_export=FAIL ({backend_error})")
        else:
            backend_used = "comtypes"
            notes.append("backend_comtypes_export=OK")

    if not any(exported):
        win32_client = None
        try:
            win32_client = importlib.import_module("win32com.client")
            notes.append("backend_win32com_import=OK")
        except Exception as error:
            notes.append(f"backend_win32com_import=FAIL ({format_exception_message(error)})")

        if win32_client is not None:
            win32_exported, backend_error = export_with_powerpoint_backend(
                win32_client,
                "win32com",
                pptx_path,
                out_dir,
                slide_count,
                export_width,
                export_height,
            )
            if any(win32_exported):
                exported = win32_exported
                backend_used = "win32com"
            if backend_error:
                notes.append(f"backend_win32com_export=FAIL ({backend_error})")
            else:
                notes.append("backend_win32com_export=OK")

    exported_count = sum(1 for path in exported if path and os.path.isfile(path))
    notes.append(f"exported_images={exported_count}/{max(slide_count, len(exported))}")
    append_pptx_export_diagnostics(
        root,
        pptx_path,
        out_dir,
        slide_count,
        export_width,
        export_height,
        notes,
    )
    if backend_used:
        print(f"INFO PPTX image export backend: {backend_used}", file=sys.stderr)
    elif exported_count == 0:
        print(
            "WARN PPTX slide image export produced 0 files. "
            "See Data/AIPresenter/Temp/pptx_export_diagnostics.txt for backend details.",
            file=sys.stderr,
        )

    return exported


def load_rows_from_single_pptx(
    root: str,
    pptx_path: str,
    start_index: int,
    deck_order: int,
    export_width: int,
    export_height: int,
) -> list[tuple[int, str, str, str]]:
    try:
        Presentation = importlib.import_module("pptx").Presentation
    except Exception:
        print("PPTX detected but python-pptx is missing. Install dependencies.", file=sys.stderr)
        return []

    presentation = Presentation(pptx_path)
    slide_count = len(presentation.slides)
    exported_abs_paths = export_pptx_slides_to_images(
        root,
        pptx_path,
        slide_count,
        deck_order,
        export_width,
        export_height,
    )

    rows: list[tuple[int, str, str, str]] = []
    base_name = os.path.basename(pptx_path)
    for local_idx, slide in enumerate(presentation.slides, start=1):
        global_idx = start_index + (local_idx - 1)
        title = ""
        try:
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip()
        except Exception:
            title = ""
        if not title:
            title = f"Slide {local_idx}"

        abs_path = ""
        rel_path = f"{base_name}#slide={local_idx}"
        if local_idx - 1 < len(exported_abs_paths):
            candidate = exported_abs_paths[local_idx - 1]
            if candidate and os.path.isfile(candidate):
                rel_candidate = os.path.normpath(os.path.relpath(candidate, root))
                abs_path = rel_candidate
                rel_path = rel_candidate
        rows.append((global_idx, title, abs_path, rel_path))

    if not any(r[2] for r in rows):
        print(
            "WARN PPTX text was parsed, but slide images were not exported. "
            "Audio/Q&A will still work; image display may be skipped.",
            file=sys.stderr,
        )
    return rows


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument(
        "--root",
        default=os.getcwd(),
        help="Project root (UiPath Environment.CurrentDirectory)",
    )
    p.add_argument(
        "--slides-xlsx",
        default=os.path.join("Data", "AIPresenter", "Input", "slides.xlsx"),
        help="Path to slides.xlsx (relative to --root if not absolute)",
    )
    p.add_argument(
        "--csv",
        default="",
        help="Deprecated. CSV source is no longer supported. Use --slides-xlsx.",
    )
    args = p.parse_args()
    root = os.path.normpath(args.root)

    input_dir = os.path.join(root, "Data", "AIPresenter", "Input")
    temp_dir = os.path.join(root, "Data", "AIPresenter", "Temp")
    os.makedirs(input_dir, exist_ok=True)
    os.makedirs(temp_dir, exist_ok=True)
    clear_pptx_export_root(root)

    slides_txt = os.path.join(input_dir, "slides.txt")
    gpt_slide_content = os.path.join(input_dir, "gpt_slide_content.txt")
    gpt_slide_sidecar = os.path.join(temp_dir, "gpt_slide.txt")
    for path in (slides_txt, gpt_slide_content, gpt_slide_sidecar):
        if not os.path.isfile(path):
            open(path, "w", encoding="utf-8").close()

    slides_xlsx_path = args.slides_xlsx
    if args.csv.strip():
        legacy_path = args.csv.strip()
        if legacy_path.lower().endswith(".csv"):
            detail = (
                "CSV slide source is no longer supported. "
                "Please convert to Data/AIPresenter/Input/slides.xlsx."
            )
            write_pptx_source_error(root, detail)
            print(detail, file=sys.stderr)
            sys.exit(1)
        slides_xlsx_path = legacy_path
    if not os.path.isabs(slides_xlsx_path):
        slides_xlsx_path = os.path.normpath(os.path.join(root, slides_xlsx_path))
    if slides_xlsx_path.lower().endswith(".csv"):
        detail = (
            "CSV slide source is no longer supported. "
            "Please use an Excel workbook (.xlsx) for slides input."
        )
        write_pptx_source_error(root, detail)
        print(detail, file=sys.stderr)
        sys.exit(1)

    out_tsv = os.path.join(temp_dir, "slides_iterate.tsv")

    cfg_path = os.path.join(root, "Data", "AIPresenter", "Config", "ai_settings.json")
    cfg: dict = {}
    configured_pptx_setting = ""
    if os.path.isfile(cfg_path):
        try:
            with open(cfg_path, encoding="utf-8-sig") as f:
                cfg = json.load(f)
            configured_pptx_setting = (
                (cfg.get("presentation_pptx_path") or "").strip()
                or (cfg.get("pptx_input_path") or "").strip()
                or (cfg.get("slide_pptx_path") or "").strip()
            )
        except Exception:
            cfg = {}
            configured_pptx_setting = ""

    export_width = parse_export_dimension(cfg.get("pptx_export_width"), DEFAULT_EXPORT_WIDTH)
    export_height = parse_export_dimension(cfg.get("pptx_export_height"), DEFAULT_EXPORT_HEIGHT)

    clear_pptx_source_error(root)
    clear_pptx_export_diagnostics(root)
    try:
        pptx_paths = resolve_pptx_paths(root, cfg)
    except RuntimeError as e:
        detail = str(e).strip()
        write_pptx_source_error(root, detail)
        print(detail, file=sys.stderr)
        sys.exit(1)

    if configured_pptx_setting and not pptx_paths:
        detail = (
            "Configured PPTX path not found or unreadable. "
            "Check presentation_pptx_path/pptx_input_path/slide_pptx_path."
        )
        write_pptx_source_error(root, detail)
        print(detail, file=sys.stderr)
        sys.exit(1)

    rows: list[tuple[int, str, str, str]] = []
    source_mode = "none"

    if pptx_paths:
        next_index = 1
        for deck_order, pptx_path in enumerate(pptx_paths, start=1):
            deck_rows = load_rows_from_single_pptx(
                root,
                pptx_path,
                next_index,
                deck_order,
                export_width,
                export_height,
            )
            rows.extend(deck_rows)
            next_index += len(deck_rows)

        source_mode = "pptx" if len(pptx_paths) == 1 else f"pptx-batch:{len(pptx_paths)}"
        print(
            "INFO pptx order: " + ", ".join(os.path.basename(path) for path in pptx_paths),
            file=sys.stderr,
        )
        print(
            f"INFO pptx export resolution: {export_width}x{export_height}",
            file=sys.stderr,
        )
    else:
        rows = load_rows_from_excel(root, slides_xlsx_path)
        if rows:
            source_mode = "slides.xlsx"

    if source_mode.startswith("pptx") and rows:
        exported_images = 0
        for _, _, image_path, _ in rows:
            if not image_path:
                continue
            candidate = (
                image_path
                if os.path.isabs(image_path)
                else os.path.normpath(os.path.join(root, image_path))
            )
            if os.path.isfile(candidate):
                exported_images += 1
        if exported_images == 0:
            detail = (
                "PPTX text was extracted but 0 slide images were exported. "
                "Failing fast to prevent audio-only playback. "
                "Verify Microsoft PowerPoint Desktop, Python package 'comtypes', and write permission "
                "for Data/AIPresenter/Temp/pptx_slides. "
                "See Data/AIPresenter/Temp/pptx_export_diagnostics.txt for backend diagnostics."
            )
            write_pptx_source_error(root, detail)
            print(detail, file=sys.stderr)
            sys.exit(3)
        if exported_images < len(rows):
            print(
                f"WARN Only {exported_images}/{len(rows)} slide images were exported. "
                "Slides without images may play narration audio only.",
                file=sys.stderr,
            )

    with open(out_tsv, "w", encoding="utf-8") as out:
        for idx, title, abs_path, rel in rows:
            out.write(
                f"{idx}\t{sanitize_tsv_field(title)}\t"
                f"{sanitize_tsv_field(abs_path)}\t{sanitize_tsv_field(rel)}\n"
            )

    if source_mode == "none":
        detail = (
            "No PPTX or slides.xlsx source found. Ensure at least one .pptx is available "
            "(Data/AIPresenter/Input/pptx_file_here or configured path), "
            "or provide Data/AIPresenter/Input/slides.xlsx."
        )
        write_pptx_source_error(root, detail)
        print(detail, file=sys.stderr)
        sys.exit(1)

    if source_mode.startswith("pptx") and not rows:
        detail = (
            "PPTX source was detected but no slides were extracted. "
            "Check PPTX readability and required dependencies (python-pptx/comtypes/PowerPoint)."
        )
        write_pptx_source_error(root, detail)
        print(detail, file=sys.stderr)
        sys.exit(2)

    print(f"OK source={source_mode} {len(rows)} slide(s) -> {out_tsv}")
    sys.exit(0)


if __name__ == "__main__":
    main()
