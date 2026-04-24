"""
Build composed narration from slide sources.

Priority:
1) Configured .pptx path in Data/AIPresenter/Config/ai_settings.json
2) Ordered .pptx list in Data/AIPresenter/Input/pptx_run_order.xlsx
    or Data/AIPresenter/Input/pptx_run_order_template.xlsx
3) First .pptx in Data/AIPresenter/Input/pptx_file_here
4) slides.xlsx + slide images (legacy fallback)

Writes:
- Data/AIPresenter/Temp/extracted_slides_raw.txt
- Data/AIPresenter/Input/gpt_slide_content.txt

Run from project folder:
    py Data\\AIPresenter\\Scripts\\pipeline_extract_and_compose.py --root "%CD%"
"""
from __future__ import annotations

import argparse
import base64
import json
import os
import re
import sys
import tempfile
import urllib.error
import urllib.parse
import urllib.request
from typing import Any

RUN_ORDER_DEFAULT = os.path.join("Data", "AIPresenter", "Input", "pptx_run_order.xlsx")
RUN_ORDER_TEMPLATE_XLSX_DEFAULT = os.path.join(
    "Data", "AIPresenter", "Input", "pptx_run_order_template.xlsx"
)
NARRATION_MODE_XLSX_DEFAULT = os.path.join(
    "Data", "AIPresenter", "Input", "slide_narration_mode.xlsx"
)
SLIDE_INDEX_MAP_XLSX_DEFAULT = os.path.join(
    "Data", "AIPresenter", "Input", "slide_index_mapping.xlsx"
)
SUPPORTED_WORKBOOK_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
SLIDES_ITERATE_TSV_DEFAULT = os.path.join(
    "Data", "AIPresenter", "Temp", "slides_iterate.tsv"
)
PPTX_SOURCE_ERROR_FILE = os.path.join(
    "Data", "AIPresenter", "Temp", "pptx_source_resolution_error.txt"
)
DEFAULT_OCR_PASSES = 2
MAX_OCR_PASSES = 3
READING_ORDER_ROW_MARGIN_EMU = 40000
ILLEGAL_XLSX_CHAR_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def configure_stdio_utf8_fallback() -> None:
    for stream_name in ("stdout", "stderr"):
        stream = getattr(sys, stream_name, None)
        if stream is None:
            continue
        reconfigure = getattr(stream, "reconfigure", None)
        if not callable(reconfigure):
            continue
        try:
            reconfigure(encoding="utf-8", errors="backslashreplace")
        except Exception:
            pass


configure_stdio_utf8_fallback()


def load_json(path: str) -> dict:
    with open(path, encoding="utf-8-sig") as f:
        return json.load(f)


def http_post_json(url: str, headers: dict, payload: dict) -> dict:
    data = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(url, data=data, headers=headers, method="POST")
    with urllib.request.urlopen(req, timeout=300) as resp:
        return json.loads(resp.read().decode("utf-8"))


def parse_ocr_passes(value: Any, default_value: int = DEFAULT_OCR_PASSES) -> int:
    try:
        parsed = int(float(str(value).strip()))
    except (TypeError, ValueError):
        parsed = default_value
    if parsed < 1:
        return 1
    if parsed > MAX_OCR_PASSES:
        return MAX_OCR_PASSES
    return parsed


def sanitize_excel_cell_text(value: Any) -> str:
    return ILLEGAL_XLSX_CHAR_RE.sub(" ", str(value or ""))


def _encode_image_for_vision(image_path: str) -> tuple[str, str]:
    with open(image_path, "rb") as f:
        b64 = base64.b64encode(f.read()).decode("ascii")
    ext = os.path.splitext(image_path)[1].lower()
    mime = "image/jpeg" if ext in (".jpg", ".jpeg") else "image/png" if ext == ".png" else "image/jpeg"
    return mime, b64


def _vision_extract_single_pass(
    api_key: str,
    model: str,
    mime: str,
    b64: str,
    slide_label: str,
    prompt_text: str,
    max_tokens: int,
) -> str:
    user_text = (
        f"Slide: {slide_label}\n"
        f"{prompt_text}"
    )
    payload = {
        "model": model,
        "messages": [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": user_text},
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:{mime};base64,{b64}",
                            "detail": "high",
                        },
                    },
                ],
            }
        ],
        "max_tokens": max_tokens,
        "temperature": 0.2,
    }
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    out = http_post_json("https://api.openai.com/v1/chat/completions", headers, payload)
    return out["choices"][0]["message"]["content"].strip()


def _ocr_text_to_lines(text: str) -> list[str]:
    lines: list[str] = []
    for raw in str(text or "").splitlines():
        line = " ".join(raw.split()).strip()
        if not line:
            continue
        if _is_no_text_ocr_result(line):
            continue
        if line.lower().startswith("slide:"):
            continue
        lines.append(line)
    return lines


def _merge_ocr_pass_outputs(pass_outputs: list[str]) -> str:
    merged: list[str] = []
    seen: set[str] = set()

    for text in pass_outputs:
        if _is_no_text_ocr_result(text):
            continue
        for line in _ocr_text_to_lines(text):
            key = _normalize_for_match(line)
            if not key or key in seen:
                continue
            seen.add(key)
            merged.append(line)

    if not merged:
        return "[No readable text detected]"

    return "\n".join(merged)


def _cleanup_temp_paths(paths: list[str]) -> None:
    for path in paths:
        try:
            if path and os.path.isfile(path):
                os.remove(path)
        except OSError:
            pass


def _build_ocr_retry_tiles(image_path: str) -> list[str]:
    """Build enlarged overlapping tiles to improve tiny-text OCR on dense slides."""
    try:
        from PIL import Image
    except Exception:
        return []

    tile_paths: list[str] = []
    try:
        with Image.open(image_path) as img:
            rgb = img.convert("RGB")
            width, height = rgb.size
            if width < 800 or height < 450:
                return []

            tile_w = max(int(width * 0.62), 1)
            tile_h = max(int(height * 0.62), 1)
            origins = [
                (0, 0),
                (max(width - tile_w, 0), 0),
                (0, max(height - tile_h, 0)),
                (max(width - tile_w, 0), max(height - tile_h, 0)),
                (max((width - tile_w) // 2, 0), max((height - tile_h) // 2, 0)),
            ]

            unique_origins: list[tuple[int, int]] = []
            seen: set[tuple[int, int]] = set()
            for ox, oy in origins:
                key = (ox, oy)
                if key in seen:
                    continue
                seen.add(key)
                unique_origins.append(key)

            resampling = getattr(Image, "Resampling", Image)
            lanczos = getattr(resampling, "LANCZOS", getattr(Image, "LANCZOS", 1))

            for ox, oy in unique_origins:
                right = min(width, ox + tile_w)
                bottom = min(height, oy + tile_h)
                tile = rgb.crop((ox, oy, right, bottom))
                if tile.width <= 0 or tile.height <= 0:
                    continue

                enlarged = tile.resize((max(tile.width * 2, tile.width), max(tile.height * 2, tile.height)), lanczos)
                fd, tmp_path = tempfile.mkstemp(prefix="ocr_tile_", suffix=".png")
                os.close(fd)
                enlarged.save(tmp_path, format="PNG")
                tile_paths.append(tmp_path)
    except Exception as e:
        print(f"WARN OCR tile preparation failed for {image_path}: {e}", file=sys.stderr)
        _cleanup_temp_paths(tile_paths)
        return []

    return tile_paths


def vision_extract(
    api_key: str,
    model: str,
    image_path: str,
    slide_label: str,
    max_passes: int = 1,
) -> str:
    mime, b64 = _encode_image_for_vision(image_path)
    passes = parse_ocr_passes(max_passes, default_value=1)

    strict_prompt = (
        "Extract only the visible text in reading order. "
        "Include visible URLs/hyperlinks exactly as shown. "
        "Do not summarize, interpret, or add external information. "
        "If no readable text is visible, return exactly: [No readable text detected]."
    )
    dense_prompt = (
        "Second OCR pass for dense slides: prioritize tiny text, table cells, chart labels, "
        "footnotes, and callout text when legible. "
        "Return only visible text, one item per line. "
        "Include visible URLs/hyperlinks exactly as shown. "
        "Do not summarize, interpret, or add external information. "
        "If no readable text is visible, return exactly: [No readable text detected]."
    )
    org_chart_prompt = (
        "Third OCR pass for very dense slides (org charts, names, role titles, and bullet lists). "
        "Extract all visible names, titles, labels, and bullet text in reading order. "
        "Return only visible text. Do not summarize. "
        "If no readable text is visible, return exactly: [No readable text detected]."
    )

    outputs: list[str] = [
        _vision_extract_single_pass(
            api_key=api_key,
            model=model,
            mime=mime,
            b64=b64,
            slide_label=slide_label,
            prompt_text=strict_prompt,
            max_tokens=4096,
        )
    ]

    if passes >= 2:
        outputs.append(
            _vision_extract_single_pass(
                api_key=api_key,
                model=model,
                mime=mime,
                b64=b64,
                slide_label=slide_label,
                prompt_text=dense_prompt,
                max_tokens=6144,
            )
        )

    if passes >= 3:
        outputs.append(
            _vision_extract_single_pass(
                api_key=api_key,
                model=model,
                mime=mime,
                b64=b64,
                slide_label=slide_label,
                prompt_text=org_chart_prompt,
                max_tokens=6144,
            )
        )

    return _merge_ocr_pass_outputs(outputs)


def vision_extract_with_tiling_retry(
    api_key: str,
    model: str,
    image_path: str,
    slide_label: str,
    max_passes: int = 1,
) -> str:
    primary = vision_extract(
        api_key=api_key,
        model=model,
        image_path=image_path,
        slide_label=slide_label,
        max_passes=max_passes,
    )
    if not _is_no_text_ocr_result(primary):
        return primary

    tile_paths = _build_ocr_retry_tiles(image_path)
    if not tile_paths:
        return primary

    tile_prompt = (
        "OCR retry for a zoomed tile from a dense slide. "
        "Extract only visible text (names, titles, labels, bullets) in reading order. "
        "Do not summarize. If no readable text is visible, return exactly: [No readable text detected]."
    )

    outputs: list[str] = [primary]
    try:
        for tile_idx, tile_path in enumerate(tile_paths, start=1):
            try:
                mime, b64 = _encode_image_for_vision(tile_path)
                outputs.append(
                    _vision_extract_single_pass(
                        api_key=api_key,
                        model=model,
                        mime=mime,
                        b64=b64,
                        slide_label=f"{slide_label} [tile {tile_idx}]",
                        prompt_text=tile_prompt,
                        max_tokens=4096,
                    )
                )
            except Exception as e:
                print(f"WARN OCR tile pass failed for {slide_label} tile {tile_idx}: {e}", file=sys.stderr)
    finally:
        _cleanup_temp_paths(tile_paths)

    merged = _merge_ocr_pass_outputs(outputs)
    if not _is_no_text_ocr_result(merged):
        print(f"INFO OCR tile retry recovered text for slide {slide_label}.", file=sys.stderr)
        return merged
    return primary


def chat_completion(api_key: str, model: str, system: str, user: str) -> str:
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
        "temperature": 0.2,
        "max_tokens": 8192,
    }
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    out = http_post_json("https://api.openai.com/v1/chat/completions", headers, payload)
    return out["choices"][0]["message"]["content"].strip()


def parse_slides_blocks(gpt_text: str) -> list[tuple[int, str, str]]:
    """Return [(index, title, body), ...] from === SLIDE | ... === blocks."""
    pattern = re.compile(
        r"===\s*SLIDE\s*\|\s*(\d+)\s*\|\s*([^=]+?)\s*===\s*(.*?)(?=\n===\s*SLIDE|\Z)",
        re.DOTALL | re.IGNORECASE,
    )
    found = []
    for m in pattern.finditer(gpt_text):
        idx = int(m.group(1))
        title = m.group(2).strip()
        body = m.group(3).strip()
        found.append((idx, title, body))
    return sorted(found, key=lambda x: x[0])


def parse_raw_slide_blocks(raw_text: str) -> dict[int, tuple[str, str]]:
    """Return {index: (title, raw_body)} from === SLIDE n (title) === blocks."""
    pattern = re.compile(
        r"===\s*SLIDE\s*(\d+)\s*\((.*?)\)\s*===\s*(.*?)(?=\n===\s*SLIDE\s*\d+\s*\(|\Z)",
        re.DOTALL | re.IGNORECASE,
    )
    out: dict[int, tuple[str, str]] = {}
    for m in pattern.finditer(raw_text):
        idx = int(m.group(1))
        title = (m.group(2) or "").strip()
        body = (m.group(3) or "").strip()
        out[idx] = (title, body)
    return out


URL_TOKEN_RE = re.compile(r"(?i)\b(?:https?://|www\.)[^\s<>()\[\]{}\"']+")


def _strip_url_tail(token: str) -> str:
    return str(token or "").rstrip(".,;:!?)]}\"'")


def _extract_urls(text: str) -> list[str]:
    urls: list[str] = []
    seen: set[str] = set()
    for match in URL_TOKEN_RE.finditer(str(text or "")):
        token = _strip_url_tail(match.group(0).strip())
        if not token:
            continue
        key = token.lower()
        if key in seen:
            continue
        seen.add(key)
        urls.append(token)
    return urls


def _domain_from_url(token: str) -> str:
    candidate = str(token or "").strip()
    if not candidate:
        return ""

    if not re.match(r"(?i)^[a-z][a-z0-9+.-]*://", candidate):
        candidate = "https://" + candidate

    try:
        parsed = urllib.parse.urlparse(candidate)
    except Exception:
        return ""

    host = (parsed.netloc or parsed.path or "").strip().lower()
    if not host:
        return ""
    if "@" in host:
        host = host.split("@", 1)[1]
    if ":" in host:
        host = host.split(":", 1)[0]
    host = host.strip(".")
    if host.startswith("www."):
        host = host[4:]
    return host


def _is_government_domain(domain: str) -> bool:
    token = str(domain or "").strip().lower()
    if not token:
        return False
    if token.endswith(".gov") or ".gov." in token:
        return True
    if token.endswith(".gob") or ".gob." in token:
        return True
    if token.endswith(".gouv") or ".gouv." in token:
        return True
    if re.search(r"(^|\.)go\.[a-z]{2,}$", token):
        return True
    return False


def _friendly_site_label(domain: str) -> str:
    token = str(domain or "").strip().lower()
    if not token:
        return "the referenced website"

    primary = token.split(".", 1)[0]
    primary = re.sub(r"[^a-z0-9]+", " ", primary, flags=re.IGNORECASE).strip()
    if primary and primary not in {"www", "com", "org", "net", "gov"}:
        return f"the {primary.title()} website"
    return f"the {token} website"


def _strip_hyperlink_lines(text: str) -> str:
    cleaned = str(text or "")
    # Drop standalone "Hyperlink: ..." lines and inline label remnants.
    cleaned = re.sub(r"(?im)^\s*hyperlink\s*:\s*.*(?:\r?\n)?", "", cleaned)
    cleaned = re.sub(r"(?i)\bhyperlink\s*:\s*", "", cleaned)
    cleaned = re.sub(r"[ \t]{2,}", " ", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def _replace_urls_with_site_labels(text: str) -> str:
    def _replace(match: re.Match[str]) -> str:
        token = _strip_url_tail(match.group(0))
        domain = _domain_from_url(token)
        return _friendly_site_label(domain)

    replaced = URL_TOKEN_RE.sub(_replace, str(text or ""))
    replaced = re.sub(r"[ \t]{2,}", " ", replaced)
    replaced = re.sub(r"\n{3,}", "\n\n", replaced)
    return replaced.strip()


def _build_link_summary(raw_body: str) -> str:
    urls = _extract_urls(raw_body)
    if not urls:
        return ""

    domains: list[str] = []
    seen: set[str] = set()
    for url in urls:
        domain = _domain_from_url(url)
        if not domain:
            continue
        if domain in seen:
            continue
        seen.add(domain)
        domains.append(domain)

    if any(_is_government_domain(domain) for domain in domains):
        return "This link points to an official government website and related official information."
    if domains:
        return (
            f"This link appears to direct viewers to {_friendly_site_label(domains[0])} "
            "for additional information related to this slide."
        )
    return "This slide includes a reference link for additional information."


def enforce_hyperlink_summary_coverage(
    composed_text: str,
    raw_blocks: dict[int, tuple[str, str]],
) -> str:
    """Ensure narration does not read raw URLs and adds short hyperlink purpose summaries."""
    parsed = parse_slides_blocks(composed_text)
    if not parsed:
        return composed_text

    out_blocks: list[str] = []
    for idx, composed_title, composed_body in parsed:
        raw_title, raw_body = raw_blocks.get(idx, (composed_title, ""))
        title = _sanitize_slide_title(composed_title or raw_title or f"Slide {idx}")

        body = _strip_hyperlink_lines(composed_body)
        body = _replace_urls_with_site_labels(body)

        link_summary = _build_link_summary(raw_body)
        if link_summary:
            lowered = body.lower()
            has_link_summary = (
                "this link" in lowered
                or "official government website" in lowered
                or "reference link" in lowered
            )
            if not has_link_summary:
                if body and not re.search(r"[.!?]\s*$", body):
                    body += "."
                body = (body + " " + link_summary).strip()

        if not body:
            body = "This slide is primarily visual, and no readable text was extracted."

        out_blocks.append(f"=== SLIDE | {idx} | {title} ===\n{body}".rstrip())

    return "\n\n".join(out_blocks).strip()


def _normalize_for_match(text: str) -> str:
    lowered = str(text or "").lower()
    lowered = re.sub(r"[^a-z0-9]+", " ", lowered)
    return " ".join(lowered.split()).strip()


UNAVAILABLE_BODY_PATTERNS = (
    "content for this slide is currently unavailable",
    "let me move to the next point",
)


def _dedupe_repeated_half(line: str) -> str:
    tokens = str(line or "").split()
    if len(tokens) >= 4 and len(tokens) % 2 == 0:
        half = len(tokens) // 2
        if tokens[:half] == tokens[half:]:
            return " ".join(tokens[:half])
    return str(line or "")


def _is_title_fragment(line_norm: str, title_norm: str) -> bool:
    if not line_norm or not title_norm:
        return False
    if line_norm == title_norm:
        return False
    if len(line_norm.split()) < 2:
        return False
    padded_title = f" {title_norm} "
    padded_line = f" {line_norm} "
    return padded_line in padded_title


def _extract_raw_lines_for_full_mode(raw_body: str, slide_title: str) -> list[str]:
    title_norm = _normalize_for_match(slide_title)
    picked: list[str] = []
    seen: set[str] = set()

    for raw_line in str(raw_body or "").splitlines():
        line = " ".join(str(raw_line).replace("\t", " ").split()).strip()
        line = _dedupe_repeated_half(line).strip()
        if not line:
            continue

        line_norm = _normalize_for_match(line)
        if not line_norm:
            continue

        if line_norm in {"slide text", "presenter notes", "internal"}:
            continue

        if line_norm == title_norm:
            continue

        if _is_title_fragment(line_norm, title_norm):
            continue

        if "no extractable text found on this slide" in line_norm:
            continue

        if "image file not found" in line_norm:
            continue

        if len(line) < 3:
            continue

        if re.fullmatch(r"\d{1,3}", line):
            continue

        if line_norm in seen:
            continue

        seen.add(line_norm)
        picked.append(line)

    return picked


def _sanitize_slide_title(title: str) -> str:
    cleaned = re.sub(r"[\t\r\n\x0b\x0c]+", " ", str(title or ""))
    cleaned = " ".join(cleaned.split())
    cleaned = cleaned.replace("=", " ")
    cleaned = " ".join(cleaned.split()).strip()
    return cleaned or "Untitled Slide"


def _contains_unavailable_placeholder(body: str) -> bool:
    lowered = str(body or "").strip().lower()
    if not lowered:
        return False
    return any(token in lowered for token in UNAVAILABLE_BODY_PATTERNS)


def _strip_unavailable_placeholder(body: str) -> str:
    text = str(body or "")
    text = re.sub(
        r"(?i)it seems the content for this slide is currently unavailable\.?\s*",
        "",
        text,
    )
    text = re.sub(r"(?i)let me move to the next point\.?\s*", "", text)
    return text.strip()


def _is_raw_content_missing(raw_body: str) -> bool:
    lowered = str(raw_body or "").lower()
    if "[no extractable text found on this slide.]" in lowered:
        return True
    if "[image file not found:" in lowered:
        return True
    return len(_extract_raw_lines_for_full_mode(raw_body, "")) == 0


def _as_sentence(line: str) -> str:
    cleaned = " ".join(str(line or "").split()).strip()
    if not cleaned:
        return ""
    if cleaned.endswith(":"):
        cleaned = cleaned.rstrip(":").rstrip()
    if re.search(r"[.!?]$", cleaned):
        return cleaned
    return cleaned + "."


def _count_sentences(text: str) -> int:
    parts = re.findall(r"[^.!?]+(?:[.!?]+|$)", str(text or ""))
    return sum(1 for part in parts if part.strip())


def _split_repeated_numbered_items(line: str) -> str:
    text = str(line or "")
    matches = list(re.finditer(r"(?<!\w)(\d+\.)\s+", text))
    if len(matches) < 2:
        return text

    parts: list[str] = []
    last = 0
    for idx, match in enumerate(matches):
        if idx == 0:
            continue
        start = match.start()
        head = text[last:start].rstrip(" ;,")
        if head:
            parts.append(head)
        last = start

    tail = text[last:].strip()
    if tail:
        parts.append(tail)

    return "; ".join(part for part in parts if part).strip() or text


def _prepare_raw_lines_for_narration(raw_lines: list[str]) -> list[str]:
    normalized: list[str] = []
    seen: set[str] = set()

    for raw_line in raw_lines:
        line = " ".join(str(raw_line or "").split()).strip()
        line = _dedupe_repeated_half(line).strip()
        line = _split_repeated_numbered_items(line).strip()
        if not line:
            continue
        key = _normalize_for_match(line)
        if not key or key in seen:
            continue
        seen.add(key)
        normalized.append(line)

    merged: list[str] = []
    continuation_starts = (
        "below ",
        "subject to ",
        "provided ",
        "including ",
        "except ",
        "otherwise ",
        "that ",
        "which ",
    )

    for line in normalized:
        if not merged:
            merged.append(line)
            continue

        current = line.strip()
        previous = merged[-1].strip()
        current_lower = current.lower()
        is_continuation = (
            current[:1].islower()
            or current.startswith("(")
            or any(current_lower.startswith(prefix) for prefix in continuation_starts)
        )

        if len(previous) <= 24 and previous.endswith(":"):
            merged[-1] = f"{previous} {current}".strip()
            continue

        if is_continuation:
            merged[-1] = f"{previous} {current}".strip()
            continue

        merged.append(current)

    return merged


def _is_list_heavy_raw_lines(raw_lines: list[str]) -> bool:
    if len(raw_lines) < 5:
        return False

    structured = 0
    for line in raw_lines:
        compact = " ".join(str(line or "").split()).strip()
        if not compact:
            continue
        if re.match(r"^(?:\d+[\.\)]|[A-Za-z]\.|[-*])\s*", compact):
            structured += 1
            continue
        if ":" in compact:
            structured += 1
            continue
        if len(compact) <= 90:
            structured += 1

    return structured >= max(4, len(raw_lines) // 2)


def _choose_detail_prefix(lines: list[str]) -> str:
    joined = " ".join(lines[:3]).lower()
    if re.search(
        r"\b(requirements?|required|documents?|claim|form|certificate|proof|statement|id|ids)\b",
        joined,
    ):
        return "Required details include"
    if re.search(
        r"\b(process|review|send|notify|screen|endorse|approve|credit|email)\b",
        joined,
    ):
        return "The process includes"
    return "Key details include"


def _compose_faithful_body(raw_lines: list[str], mode: str) -> str:
    normalized_lines = _prepare_raw_lines_for_narration(raw_lines)
    if not normalized_lines:
        return ""

    if mode == "short":
        chosen = normalized_lines[:2]
        return " ".join(part for part in (_as_sentence(line) for line in chosen) if part).strip()

    if len(normalized_lines) <= 4:
        return " ".join(
            part for part in (_as_sentence(line) for line in normalized_lines) if part
        ).strip()

    sentences: list[str] = []
    remaining = list(normalized_lines)

    first_line = normalized_lines[0]
    first_is_explanatory = (
        len(first_line) > 85
        or first_line.rstrip().endswith(":")
        or re.search(
            r"\b(following|required|eligible|benefit|provide|submit|include|must|can be|will be)\b",
            first_line,
            re.IGNORECASE,
        )
        is not None
    )

    first = _as_sentence(first_line) if first_is_explanatory else ""
    if first:
        sentences.append(first)
        remaining = normalized_lines[1:]

    if not remaining:
        return " ".join(sentences).strip()

    chunk_size = 4 if len(remaining) > 6 else 3
    first_prefix = _choose_detail_prefix(remaining)
    for offset in range(0, len(remaining), chunk_size):
        chunk = remaining[offset: offset + chunk_size]
        if not chunk:
            continue
        chunk_text = "; ".join(" ".join(item.split()) for item in chunk).strip()
        if chunk_text:
            if offset == 0:
                label = first_prefix
            else:
                label = "Additional details include"
            sentences.append(f"{label}: {chunk_text}.")

    return " ".join(sentences).strip()


def enforce_full_mode_coverage(
    composed_text: str,
    raw_blocks: dict[int, tuple[str, str]],
    mode_map: dict[int, str],
) -> str:
    """Repair placeholder outputs and keep narration tightly bound to slide raw content."""
    parsed = parse_slides_blocks(composed_text)
    composed_by_idx: dict[int, tuple[str, str]] = {
        idx: (title, body) for idx, title, body in parsed
    }

    if not raw_blocks:
        return composed_text

    out_blocks: list[str] = []
    for idx in sorted(raw_blocks.keys()):
        raw_title, raw_body = raw_blocks[idx]
        composed_title, composed_body = composed_by_idx.get(idx, (raw_title, ""))
        title = _sanitize_slide_title(composed_title or raw_title or f"Slide {idx}")
        mode = mode_map.get(idx, "full")
        raw_lines = _extract_raw_lines_for_full_mode(raw_body, raw_title)
        raw_missing = _is_raw_content_missing(raw_body)

        source_body = (composed_body or "").strip()
        body = _strip_unavailable_placeholder(source_body)
        body = re.sub(
            r"(?is)\bto ensure full coverage of this slide, the listed details are:\s*.*$",
            "",
            body,
        ).strip()

        needs_rebuild = False
        if raw_missing:
            if not body:
                body = "This slide is primarily visual, and no readable text was extracted."
        else:
            if not body:
                needs_rebuild = True
            if _contains_unavailable_placeholder(source_body):
                needs_rebuild = True

            raw_len = len(" ".join(raw_lines))
            sentence_count = _count_sentences(body)
            list_heavy = _is_list_heavy_raw_lines(raw_lines)
            if mode != "short":
                if body and raw_len > 260 and len(body) < max(120, int(raw_len * 0.33)):
                    needs_rebuild = True
                if body and raw_len > 0 and len(body) > int(raw_len * 1.8):
                    needs_rebuild = True
                if body and list_heavy:
                    min_sentences = 3 if len(raw_lines) >= 8 else 2
                    if sentence_count < min_sentences:
                        needs_rebuild = True
            elif body and len(body) > max(240, int(raw_len * 1.3)):
                needs_rebuild = True

            if needs_rebuild:
                body = _compose_faithful_body(raw_lines, mode)

        if not body:
            body = "This slide is primarily visual, and no readable text was extracted."

        out_blocks.append(f"=== SLIDE | {idx} | {title} ===\n{body}".rstrip())

    return "\n\n".join(out_blocks).strip()


def parse_narration_mode(value: Any) -> tuple[str, bool]:
    """Return (mode, is_known). mode is one of: full, short."""
    token = str(value or "").strip().lower().replace("-", "_").replace(" ", "_")
    if token in {"", "full", "detailed", "long", "normal", "default"}:
        return "full", True
    if token in {"short", "brief", "summary", "concise", "ngan", "short_summary"}:
        return "short", True
    return "full", False


def normalize_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


def normalize_pptx_key(value: Any) -> str:
    return str(value or "").strip().strip('"').replace("\\", "/").lower()


def candidate_pptx_keys(value: Any) -> list[str]:
    token = normalize_pptx_key(value)
    if not token:
        return []

    base = os.path.basename(token)
    candidates = [token]
    if base and base != token:
        candidates.append(base)
    return candidates


def build_pptx_slide_lookup(
    root: str,
    rows: list[tuple[int, str, str, str, int]],
) -> dict[tuple[str, int], int]:
    """Map (pptx_key, local_slide_index) -> global_slide_index."""
    lookup: dict[tuple[str, int], int] = {}

    for global_idx, _title, _text, pptx_path, local_idx in rows:
        keys: list[str] = []

        normalized_abs = normalize_pptx_key(pptx_path)
        if normalized_abs:
            keys.extend(candidate_pptx_keys(normalized_abs))

        try:
            rel = normalize_pptx_key(os.path.relpath(pptx_path, root))
            if rel:
                keys.extend(candidate_pptx_keys(rel))
        except ValueError:
            pass

        seen: set[str] = set()
        for key in keys:
            if not key or key in seen:
                continue
            seen.add(key)
            lookup[(key, local_idx)] = global_idx

    return lookup


def build_slide_index_source_meta(
    rows: list[tuple[int, str, str, str, int]],
) -> dict[int, tuple[str, int]]:
    meta: dict[int, tuple[str, int]] = {}
    for global_idx, _title, _text, pptx_path, local_idx in rows:
        meta[global_idx] = (os.path.basename(pptx_path), local_idx)
    return meta


def write_slide_index_mapping_xlsx(
    root: str,
    output_xlsx_path: str,
    rows: list[tuple[int, str, str, str, int]],
) -> None:
    try:
        from openpyxl import Workbook
    except Exception:
        print(
            "WARN Cannot write slide_index_mapping.xlsx because openpyxl is missing.",
            file=sys.stderr,
        )
        return

    os.makedirs(os.path.dirname(output_xlsx_path), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "slide_index_mapping"
    ws.append(
        [
            "global_slide_index",
            "pptx_file",
            "file_slide_index",
            "pptx_path",
            "slide_title",
        ]
    )

    for global_idx, title, _text, pptx_path, local_idx in rows:
        try:
            rel_path = os.path.relpath(pptx_path, root).replace("\\", "/")
        except ValueError:
            rel_path = pptx_path.replace("\\", "/")
        ws.append(
            [
                global_idx,
                sanitize_excel_cell_text(os.path.basename(pptx_path)),
                local_idx,
                sanitize_excel_cell_text(rel_path),
                sanitize_excel_cell_text(title),
            ]
        )

    try:
        wb.save(output_xlsx_path)
    finally:
        wb.close()


def read_slide_narration_mode_xlsx(
    workbook_path: str,
    pptx_slide_lookup: dict[tuple[str, int], int] | None = None,
) -> tuple[dict[int, str], list[str]]:
    """Load per-slide narration mode config from Excel workbook.

    Supported target selectors per row:
    - slide_index (global, preferred when present)
    - OR pptx_file + file_slide_index (resolved via slide_index_mapping)

    Required:
    - mode: full | short
    """
    mode_map: dict[int, str] = {}
    warnings: list[str] = []

    if not os.path.isfile(workbook_path):
        return mode_map, warnings

    try:
        from openpyxl import load_workbook
    except Exception:
        warnings.append(
            "WARN slide_narration_mode.xlsx found but openpyxl is missing. "
            "All slides will use FULL narration mode."
        )
        return mode_map, warnings

    workbook = None
    try:
        workbook = load_workbook(workbook_path, data_only=True, read_only=True)
        sheet = workbook.active
        header_row = None
        for values in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
            header_row = list(values)
            break
        if not header_row:
            warnings.append(
                "WARN slide_narration_mode.xlsx is empty. All slides will use FULL narration mode."
            )
            return mode_map, warnings

        def norm_header(value: Any) -> str:
            return normalize_header(value)

        headers = [norm_header(v) for v in header_row]
        index_aliases = {
            "slide_index",
            "index",
            "slide",
            "slide_no",
            "slide_number",
            "global_slide_index",
            "global_index",
        }
        mode_aliases = {"mode", "narration_mode", "script_mode", "style"}
        file_aliases = {
            "pptx_file",
            "pptx",
            "file",
            "file_name",
            "filename",
            "pptx_path",
            "source_file",
            "source_pptx",
        }
        local_index_aliases = {
            "file_slide_index",
            "pptx_slide_index",
            "slide_in_file",
            "local_slide_index",
            "local_index",
            "source_slide_index",
        }

        index_idx = next((i for i, h in enumerate(headers) if h in index_aliases), -1)
        mode_idx = next((i for i, h in enumerate(headers) if h in mode_aliases), -1)
        file_idx = next((i for i, h in enumerate(headers) if h in file_aliases), -1)
        local_index_idx = next((i for i, h in enumerate(headers) if h in local_index_aliases), -1)

        has_global_selector = index_idx >= 0
        has_file_selector = file_idx >= 0 and local_index_idx >= 0

        if mode_idx < 0 or (not has_global_selector and not has_file_selector):
            warnings.append(
                "WARN slide_narration_mode.xlsx missing required columns. "
                "Expected mode plus slide_index OR pptx_file + file_slide_index."
            )
            return mode_map, warnings

        for row_num, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            row = list(values)
            selector_present = any(str(v or "").strip() for v in row)
            if not selector_present:
                continue

            def row_value(idx: int) -> str:
                if idx < 0 or idx >= len(row):
                    return ""
                raw = row[idx]
                if raw is None:
                    return ""
                return str(raw).strip()

            resolved_slide_idx: int | None = None

            idx_raw = row_value(index_idx) if has_global_selector else ""
            if idx_raw:
                try:
                    resolved_slide_idx = int(float(idx_raw))
                except ValueError:
                    warnings.append(
                        f"WARN slide_narration_mode.xlsx line {row_num}: invalid slide_index '{idx_raw}', skipped."
                    )
                    continue

            if resolved_slide_idx is None and has_file_selector:
                file_raw = row_value(file_idx)
                local_raw = row_value(local_index_idx)

                if file_raw or local_raw:
                    if not file_raw or not local_raw:
                        warnings.append(
                            f"WARN slide_narration_mode.xlsx line {row_num}: "
                            "provide both pptx_file and file_slide_index."
                        )
                        continue

                    try:
                        local_idx = int(float(local_raw))
                    except ValueError:
                        warnings.append(
                            f"WARN slide_narration_mode.xlsx line {row_num}: "
                            f"invalid file_slide_index '{local_raw}', skipped."
                        )
                        continue

                    if not pptx_slide_lookup:
                        warnings.append(
                            f"WARN slide_narration_mode.xlsx line {row_num}: "
                            "cannot resolve pptx_file/file_slide_index because no PPTX source map is available."
                        )
                        continue

                    hit_idx: int | None = None
                    for key in candidate_pptx_keys(file_raw):
                        hit_idx = pptx_slide_lookup.get((key, local_idx))
                        if hit_idx is not None:
                            break

                    if hit_idx is None:
                        warnings.append(
                            f"WARN slide_narration_mode.xlsx line {row_num}: "
                            f"could not match pptx_file='{file_raw}' and file_slide_index='{local_raw}'."
                        )
                        continue

                    resolved_slide_idx = hit_idx

            if resolved_slide_idx is None:
                continue

            if resolved_slide_idx <= 0:
                warnings.append(
                    f"WARN slide_narration_mode.xlsx line {row_num}: slide index must be >= 1, skipped."
                )
                continue

            mode_raw = row_value(mode_idx)
            mode, known = parse_narration_mode(mode_raw)
            if not known:
                warnings.append(
                    f"WARN slide_narration_mode.xlsx line {row_num}: unknown mode '{mode_raw}', defaulted to full."
                )

            mode_map[resolved_slide_idx] = mode
    except Exception as e:
        warnings.append(f"WARN Could not read slide_narration_mode.xlsx: {e}")
    finally:
        try:
            if workbook is not None:
                workbook.close()
        except Exception:
            pass

    return mode_map, warnings


def build_runtime_narration_rules(mode_map: dict[int, str]) -> str:
    short_indices = sorted(idx for idx, mode in mode_map.items() if mode == "short")

    lines = [
        "RUNTIME OVERRIDE (STRICT, HIGH PRIORITY):",
        "- Keep slide raw content as the primary source of facts.",
        "- You may add up to 1-2 short contextual sentences per slide to sound human and natural.",
        "- Any added context must be generic and relevant, without fabricated details.",
        "- Do NOT invent policy details, numbers, dates, commitments, or examples not explicitly present.",
        "- Keep narration concise and professional; avoid excessive promotional expansion.",
        "- Keep speech natural like a normal presenter, with varied but simple connective phrasing.",
        "- Vary sentence openings and list lead-ins across nearby sentences and slides.",
        "- Avoid repeated scaffolding phrases such as 'it further includes', 'the slide also lists', or similar repeated list lead-ins; prefer direct continuation or a fresh sentence when extending a list.",
        "- Never output generic fallback lines like 'content is currently unavailable' unless raw block explicitly says no extractable content.",
    ]

    if short_indices:
        lines.append(
            "- SHORT narration mode applies to these slide indices: "
            + ", ".join(str(i) for i in short_indices)
            + "."
        )
        lines.extend(
            [
                "- For SHORT slides: output only 1-2 concise presenter sentences (target <= 35 words total).",
                "- SHORT slides should naturally hand over to the human presenter.",
            ]
        )
    else:
        lines.append("- No SHORT slide is configured. Use FULL mode for all slides.")

    lines.extend(
        [
            "- For slides not configured as SHORT: use FULL mode by default.",
            "- For FULL mode: cover major bullet/list entries but stay concise and factual.",
            "- For long exclusion/requirement lists in FULL mode: keep all listed items; optional context must be brief and must not replace listed details.",
            "- Keep the required === SLIDE | index | title === output format exactly.",
        ]
    )
    return "\n".join(lines)


def build_narration_mode_user_block(
    mode_map: dict[int, str],
    slide_source_meta: dict[int, tuple[str, int]] | None = None,
) -> str:
    if not mode_map:
        return (
            "NARRATION MODE CONFIG:\n"
            "- No per-slide config found. Treat all slides as FULL mode."
        )

    lines = ["NARRATION MODE CONFIG (from slide_narration_mode.xlsx):"]
    for slide_idx in sorted(mode_map):
        source_text = ""
        if slide_source_meta and slide_idx in slide_source_meta:
            pptx_name, local_idx = slide_source_meta[slide_idx]
            source_text = f" ({pptx_name} | file slide {local_idx})"
        lines.append(f"- Slide {slide_idx}{source_text}: {mode_map[slide_idx].upper()}")
    lines.append("- Any slide not listed above is FULL mode.")
    return "\n".join(lines)


def find_first_pptx(folder: str) -> str:
    if not os.path.isdir(folder):
        return ""
    for name in sorted(os.listdir(folder)):
        low = name.lower()
        if low.endswith(".pptx") and not low.startswith("~$"):
            candidate = os.path.normpath(os.path.join(folder, name))
            if os.path.isfile(candidate):
                return candidate
    return ""


def list_pptx_files(folder: str) -> list[str]:
    files: list[str] = []
    if not os.path.isdir(folder):
        return files
    for name in sorted(os.listdir(folder)):
        low = name.lower()
        if low.endswith(".pptx") and not low.startswith("~$"):
            candidate = os.path.normpath(os.path.join(folder, name))
            if os.path.isfile(candidate):
                files.append(candidate)
    return files


def format_path_for_message(root: str, path: str) -> str:
    try:
        rel = os.path.relpath(path, root)
        return rel.replace("\\", "/")
    except ValueError:
        return path.replace("\\", "/")


def build_run_order_mismatch_message(
    root: str,
    queue_path: str,
    unresolved_entries: list[str],
    folder_pptx_files: list[str],
) -> str:
    unresolved_unique: list[str] = []
    seen: set[str] = set()
    for raw in unresolved_entries:
        key = raw.strip().lower()
        if not key or key in seen:
            continue
        seen.add(key)
        unresolved_unique.append(raw.strip())

    lines: list[str] = [
        "PPTX source mismatch: run-order file and PPTX folder are not aligned.",
        f"Run-order file: {format_path_for_message(root, queue_path)}",
        "",
        "Entries listed in run-order but not found on disk:",
    ]

    if unresolved_unique:
        for raw in unresolved_unique[:30]:
            lines.append(f"- {raw}")
        if len(unresolved_unique) > 30:
            lines.append(f"- ... and {len(unresolved_unique) - 30} more")
    else:
        lines.append("- (none)")

    lines.extend(["", "PPTX files currently found in Data/AIPresenter/Input/pptx_file_here:"])
    if folder_pptx_files:
        for path in folder_pptx_files[:30]:
            lines.append(f"- {format_path_for_message(root, path)}")
        if len(folder_pptx_files) > 30:
            lines.append(f"- ... and {len(folder_pptx_files) - 30} more")
    else:
        lines.append("- (none)")

    lines.extend(
        [
            "",
            "Fix:",
            "- Update pptx_run_order_template.xlsx (or pptx_run_order.xlsx).",
            "- Ensure each pptx_path exactly matches an existing .pptx file name/path.",
        ]
    )
    return "\n".join(lines)


def clear_pptx_source_error(root: str) -> None:
    path = os.path.join(root, PPTX_SOURCE_ERROR_FILE)
    try:
        if os.path.isfile(path):
            os.remove(path)
    except OSError:
        pass


def write_pptx_source_error(root: str, message: str) -> None:
    path = os.path.join(root, PPTX_SOURCE_ERROR_FILE)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write((message or "").strip() + "\n")


def parse_enabled(value: Any) -> bool:
    if value is None:
        return True
    text = str(value).strip().lower()
    if text == "":
        return True
    return text not in {"0", "false", "no", "n", "off", "disabled"}


def resolve_pptx_candidate(root: str, input_dir: str, raw_path: str) -> str:
    value = (raw_path or "").strip().strip('"')
    if not value:
        return ""

    candidates: list[str] = []
    if os.path.isabs(value):
        candidates.append(value)
    else:
        candidates.append(os.path.join(root, value))
        candidates.append(os.path.join(input_dir, "pptx_file_here", value))

    for candidate in candidates:
        normalized = os.path.normpath(candidate)
        base_name = os.path.basename(normalized).lower()
        if base_name.startswith("~$") or not base_name.endswith(".pptx"):
            continue
        if os.path.isfile(normalized):
            return normalized
    return ""


def read_pptx_run_order(root: str, input_dir: str, queue_path: str) -> tuple[list[str], int, list[str]]:
    if not os.path.isfile(queue_path):
        return [], 0, []

    try:
        from openpyxl import load_workbook
    except Exception:
        print(
            "WARN pptx_run_order.xlsx found but openpyxl is missing. "
            "Install dependencies to use ordered multi-PPTX mode.",
            file=sys.stderr,
        )
        return [], 0, []

    try:
        workbook = load_workbook(queue_path, data_only=True, read_only=True)
        sheet = workbook.active
    except Exception as e:
        print(f"WARN Could not read PPTX run-order file: {queue_path} ({e})", file=sys.stderr)
        return [], 0, []

    header_row = None
    for values in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
        header_row = list(values)
        break

    if not header_row:
        return [], 0, []

    def norm_header(value: Any) -> str:
        return str(value or "").strip().lower().replace(" ", "_")

    headers = [norm_header(v) for v in header_row]
    path_aliases = {"pptx_path", "file_path", "path", "pptx", "file", "filename", "pptx_file"}
    order_aliases = {"run_order", "order", "stt", "index", "priority"}
    enabled_aliases = {"enabled", "is_enabled", "active", "use"}

    path_idx = next((i for i, h in enumerate(headers) if h in path_aliases), -1)
    order_idx = next((i for i, h in enumerate(headers) if h in order_aliases), -1)
    enabled_idx = next((i for i, h in enumerate(headers) if h in enabled_aliases), -1)

    if path_idx < 0:
        print(
            "WARN pptx_run_order.xlsx does not contain a path column. "
            "Expected one of: pptx_path, file_path, path, pptx.",
            file=sys.stderr,
        )
        return [], 0, []

    items: list[tuple[int, str]] = []
    unresolved_paths: list[str] = []
    rows_with_path = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = list(row)

        raw_path = ""
        if path_idx < len(values) and values[path_idx] is not None:
            raw_path = str(values[path_idx]).strip()
        if not raw_path:
            continue

        rows_with_path += 1

        enabled = True
        if enabled_idx >= 0 and enabled_idx < len(values):
            enabled = parse_enabled(values[enabled_idx])
        if not enabled:
            continue

        order_value = rows_with_path
        if order_idx >= 0 and order_idx < len(values) and values[order_idx] not in (None, ""):
            try:
                order_value = int(float(str(values[order_idx]).strip()))
            except ValueError:
                order_value = rows_with_path

        resolved = resolve_pptx_candidate(root, input_dir, raw_path)
        if not resolved:
            print(f"WARN queue path not found: {raw_path}", file=sys.stderr)
            unresolved_paths.append(raw_path)
            continue

        items.append((order_value, resolved))

    items.sort(key=lambda it: (it[0], it[1].lower()))
    ordered_paths: list[str] = []
    seen: set[str] = set()
    for _, path in items:
        key = path.lower()
        if key in seen:
            continue
        seen.add(key)
        ordered_paths.append(path)

    return ordered_paths, rows_with_path, unresolved_paths


def resolve_pptx_paths(root: str, cfg: dict) -> list[str]:
    configured = (
        (cfg.get("presentation_pptx_path") or "").strip()
        or (cfg.get("pptx_input_path") or "").strip()
        or (cfg.get("slide_pptx_path") or "").strip()
    )
    input_dir = os.path.join(root, "Data", "AIPresenter", "Input")

    if configured:
        candidate = resolve_pptx_candidate(root, input_dir, configured)
        return [candidate] if candidate else []

    queue_setting = (cfg.get("pptx_run_order_file") or "").strip()
    queue_candidates: list[str] = []
    if queue_setting:
        queue_candidates.append(
            os.path.normpath(queue_setting)
            if os.path.isabs(queue_setting)
            else os.path.normpath(os.path.join(root, queue_setting))
        )
    else:
        queue_candidates.append(os.path.join(root, RUN_ORDER_DEFAULT))
        queue_candidates.append(os.path.join(root, RUN_ORDER_TEMPLATE_XLSX_DEFAULT))

    queue_rows_detected = 0
    queue_path_with_rows = ""
    unresolved_queue_entries: list[str] = []
    for queue_path in queue_candidates:
        ext = os.path.splitext(queue_path)[1].lower()
        if ext == ".csv":
            raise RuntimeError(
                "CSV run-order is no longer supported. "
                "Use Excel workbook .xlsx for pptx_run_order_file."
            )
        if ext and ext not in SUPPORTED_WORKBOOK_EXTENSIONS:
            print(
                f"WARN Unsupported run-order extension '{ext}' for file: {queue_path}. "
                "Expected an Excel workbook (.xlsx/.xlsm/.xltx/.xltm).",
                file=sys.stderr,
            )
            continue
        queue_paths, queue_rows, unresolved = read_pptx_run_order(root, input_dir, queue_path)

        if queue_paths:
            return queue_paths
        if queue_rows > 0 and not queue_path_with_rows:
            queue_path_with_rows = queue_path
        queue_rows_detected = max(queue_rows_detected, queue_rows)
        unresolved_queue_entries.extend(unresolved)

    if queue_rows_detected > 0:
        preferred_dir = os.path.join(input_dir, "pptx_file_here")
        preferred_files = list_pptx_files(preferred_dir)

        if unresolved_queue_entries:
            queue_path_for_message = queue_path_with_rows or queue_candidates[0]
            raise RuntimeError(
                build_run_order_mismatch_message(
                    root,
                    queue_path_for_message,
                    unresolved_queue_entries,
                    preferred_files,
                )
            )

        print(
            "WARN PPTX run-order file exists but no valid entries were resolved. "
            "Falling back to folder auto-detection.",
            file=sys.stderr,
        )

    preferred_dir = os.path.join(input_dir, "pptx_file_here")
    preferred_candidate = find_first_pptx(preferred_dir)
    if preferred_candidate:
        return [preferred_candidate]

    return []


def read_slide_image_map_from_tsv(root: str) -> dict[int, str]:
    """Read slide index -> exported image path from Temp/slides_iterate.tsv."""
    path = os.path.join(root, SLIDES_ITERATE_TSV_DEFAULT)
    out: dict[int, str] = {}
    if not os.path.isfile(path):
        return out

    try:
        with open(path, encoding="utf-8-sig") as f:
            for raw_line in f:
                line = raw_line.rstrip("\r\n")
                if not line:
                    continue
                parts = line.split("\t")
                if len(parts) < 3:
                    continue

                idx_raw = (parts[0] or "").strip()
                abs_path = (parts[2] or "").strip()
                if not idx_raw or not abs_path:
                    continue

                try:
                    idx = int(float(idx_raw))
                except ValueError:
                    continue

                if idx <= 0:
                    continue

                candidate = abs_path
                if not os.path.isabs(candidate):
                    candidate = os.path.normpath(os.path.join(root, candidate))
                if os.path.isfile(candidate):
                    out[idx] = os.path.normpath(candidate)
    except Exception as e:
        print(f"WARN Could not read slides_iterate.tsv for OCR fallback: {e}", file=sys.stderr)

    return out


def _collect_shape_hyperlinks(shape) -> list[str]:
    links: list[str] = []
    seen: set[str] = set()

    def add_link(value: Any) -> None:
        token = str(value or "").strip()
        if not token:
            return
        key = token.lower()
        if key in seen:
            return
        seen.add(key)
        links.append(token)

    try:
        click_action = getattr(shape, "click_action", None)
        if click_action is not None:
            hyperlink = getattr(click_action, "hyperlink", None)
            if hyperlink is not None:
                add_link(getattr(hyperlink, "address", ""))
    except Exception:
        pass

    try:
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        run_hyperlink = getattr(run, "hyperlink", None)
                        if run_hyperlink is not None:
                            add_link(getattr(run_hyperlink, "address", ""))
                    except Exception:
                        continue
    except Exception:
        pass

    return links


def extract_shape_text(shape) -> str:
    parts: list[str] = []

    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            line = (para.text or "").strip()
            if line:
                parts.append(line)

    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            row_cells: list[str] = []
            for cell in row.cells:
                cell_text = " ".join((cell.text or "").split())
                if cell_text:
                    row_cells.append(cell_text)
            if row_cells:
                parts.append(" | ".join(row_cells))

    links = _collect_shape_hyperlinks(shape)
    if links:
        existing = "\n".join(parts).lower()
        for link in links:
            if link.lower() not in existing:
                parts.append(f"Hyperlink: {link}")

    return "\n".join(parts).strip()


def _collect_text_shape_entries(shape, source_order: int) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []

    child_shapes = getattr(shape, "shapes", None)
    if child_shapes is not None:
        child_entries: list[dict[str, Any]] = []
        try:
            for child_idx, child_shape in enumerate(child_shapes, start=1):
                child_entries.extend(
                    _collect_text_shape_entries(
                        child_shape,
                        source_order=(source_order * 1000) + child_idx,
                    )
                )
        except Exception:
            child_entries = []
        if child_entries:
            return child_entries

    text = extract_shape_text(shape)
    if not text:
        return entries

    top = int(getattr(shape, "top", 0) or 0)
    left = int(getattr(shape, "left", 0) or 0)
    height = int(getattr(shape, "height", 0) or 0)
    entries.append(
        {
            "text": text,
            "top": top,
            "left": left,
            "bottom": top + max(height, 1),
            "source_order": source_order,
        }
    )
    return entries


def _share_visual_row(row: dict[str, Any], entry: dict[str, Any]) -> bool:
    overlap = min(row["bottom"], entry["bottom"]) - max(row["top"], entry["top"])
    return overlap >= -READING_ORDER_ROW_MARGIN_EMU


def _sort_entries_by_visual_reading_order(entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not entries:
        return []

    sorted_entries = sorted(
        entries,
        key=lambda item: (item["top"], item["left"], item["source_order"]),
    )

    rows: list[dict[str, Any]] = []
    for entry in sorted_entries:
        if rows and _share_visual_row(rows[-1], entry):
            row = rows[-1]
            row["entries"].append(entry)
            row["top"] = min(row["top"], entry["top"])
            row["bottom"] = max(row["bottom"], entry["bottom"])
            continue

        rows.append(
            {
                "top": entry["top"],
                "bottom": entry["bottom"],
                "entries": [entry],
            }
        )

    ordered: list[dict[str, Any]] = []
    for row in rows:
        row["entries"].sort(
            key=lambda item: (item["left"], item["top"], item["source_order"])
        )
        ordered.extend(row["entries"])

    return ordered


def extract_slide_text_chunks(slide) -> list[str]:
    entries: list[dict[str, Any]] = []
    for source_order, shape in enumerate(slide.shapes, start=1):
        entries.extend(_collect_text_shape_entries(shape, source_order))

    ordered_entries = _sort_entries_by_visual_reading_order(entries)
    return [entry["text"] for entry in ordered_entries if entry["text"]]


def _is_no_text_ocr_result(text: str) -> bool:
    normalized = _normalize_for_match(text)
    if not normalized:
        return True
    if "no readable text detected" in normalized:
        return True
    if "no extractable text found on this slide" in normalized:
        return True
    return normalized in {
        "no readable text",
        "no text",
        "none",
        "n a",
        "na",
    }


def enrich_pptx_rows_with_image_ocr(
    api_key: str,
    vision_model: str,
    rows: list[tuple[int, str, str, str, int]],
    slide_image_map: dict[int, str],
    ocr_passes: int,
) -> list[tuple[int, str, str, str, int]]:
    """Use exported slide images as OCR fallback for PPTX slides with no extracted text."""
    if not rows or not api_key or not slide_image_map:
        return rows

    updated: list[tuple[int, str, str, str, int]] = []
    applied = 0

    for idx, title, body, pptx_path, local_idx in rows:
        body_text = str(body or "")
        has_text = bool(_extract_raw_lines_for_full_mode(body_text, title))
        has_missing_marker = "[No extractable text found on this slide.]" in body_text

        if has_text and not has_missing_marker:
            updated.append((idx, title, body_text, pptx_path, local_idx))
            continue

        image_path = slide_image_map.get(idx, "")
        if not image_path or not os.path.isfile(image_path):
            updated.append((idx, title, body_text, pptx_path, local_idx))
            continue

        try:
            ocr_text = vision_extract_with_tiling_retry(
                api_key,
                vision_model,
                image_path,
                f"{idx} - {title}",
                max_passes=ocr_passes,
            )
        except Exception as e:
            print(f"WARN OCR fallback failed for slide {idx}: {e}", file=sys.stderr)
            updated.append((idx, title, body_text, pptx_path, local_idx))
            continue

        ocr_text = (ocr_text or "").strip()
        if not ocr_text or _is_no_text_ocr_result(ocr_text):
            updated.append((idx, title, body_text, pptx_path, local_idx))
            continue

        ocr_lines = [line.strip() for line in ocr_text.splitlines() if line.strip()]
        filtered_lines = [l for l in ocr_lines if not re.search(r"slide text", l, re.IGNORECASE)]
        main_line = max(filtered_lines, key=len) if filtered_lines else (max(ocr_lines, key=len) if ocr_lines else "")
        applied += 1
        updated.append(
            (
                idx,
                title,
                main_line,
                pptx_path,
                local_idx,
            )
        )

    if applied:
        print(
            f"INFO OCR fallback populated text for {applied} slide(s) from exported slide images.",
            file=sys.stderr,
        )
    return updated


def load_slides_from_pptx(pptx_path: str, start_index: int = 1) -> list[tuple[int, str, str, str, int]]:
    try:
        from pptx import Presentation
    except Exception as e:
        raise RuntimeError(
            "python-pptx is required for PPTX mode. Install dependencies and retry."
        ) from e

    presentation = Presentation(pptx_path)
    rows: list[tuple[int, str, str, str, int]] = []
    for local_idx, slide in enumerate(presentation.slides, start=1):
        idx = start_index + (local_idx - 1)
        title = ""
        try:
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip()
        except Exception:
            title = ""
        if not title:
            title = f"Slide {local_idx}"

        text_chunks = extract_slide_text_chunks(slide)

        notes_text = ""
        try:
            if slide.has_notes_slide and slide.notes_slide is not None:
                frame = slide.notes_slide.notes_text_frame
                if frame is not None:
                    notes_text = (frame.text or "").strip()
        except Exception:
            notes_text = ""

        body_parts: list[str] = []
        if text_chunks:
            body_parts.append("\n".join(text_chunks))
        if notes_text:
            body_parts.append("Presenter notes:\n" + notes_text)
        if not body_parts:
            body_parts.append("[No extractable text found on this slide.]")

        rows.append((idx, title, "\n\n".join(body_parts).strip(), pptx_path, local_idx))
    return rows


def load_rows_from_xlsx(root: str, workbook_path: str) -> list[tuple[int, str, str]]:
    rows: list[tuple[int, str, str]] = []
    if not os.path.isfile(workbook_path):
        return rows

    try:
        from openpyxl import load_workbook
    except Exception:
        print(
            "WARN slides.xlsx found but openpyxl is missing. "
            "Install dependencies to use Excel slide source mode.",
            file=sys.stderr,
        )
        return rows

    workbook = None
    try:
        workbook = load_workbook(workbook_path, data_only=True, read_only=True)
        sheet = workbook.active
        header_row = None
        for values in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
            header_row = list(values)
            break
        if not header_row:
            return rows

        def norm_header(value: Any) -> str:
            return str(value or "").strip().lower().replace(" ", "_")

        headers = [norm_header(v) for v in header_row]
        idx_aliases = {"slide_index", "index", "slide", "slide_no", "slide_number"}
        title_aliases = {"slide_title", "title", "slide_name", "name"}
        path_aliases = {"slide_path", "path", "image_path", "file_path"}

        idx_col = next((i for i, h in enumerate(headers) if h in idx_aliases), -1)
        title_col = next((i for i, h in enumerate(headers) if h in title_aliases), -1)
        path_col = next((i for i, h in enumerate(headers) if h in path_aliases), -1)

        if idx_col < 0 or path_col < 0:
            print(
                "WARN slides.xlsx missing required columns. "
                "Expected at least: slide_index and slide_path.",
                file=sys.stderr,
            )
            return rows

        for values in sheet.iter_rows(min_row=2, values_only=True):
            row = list(values)
            idx_raw = ""
            title_raw = ""
            rel_raw = ""
            if idx_col < len(row) and row[idx_col] is not None:
                idx_raw = str(row[idx_col]).strip()
            if title_col >= 0 and title_col < len(row) and row[title_col] is not None:
                title_raw = str(row[title_col]).strip()
            if path_col < len(row) and row[path_col] is not None:
                rel_raw = str(row[path_col]).strip().strip('"')
            if not idx_raw or not rel_raw:
                continue
            try:
                idx = int(float(idx_raw))
            except ValueError:
                continue
            abs_img = rel_raw if os.path.isabs(rel_raw) else os.path.normpath(os.path.join(root, rel_raw))
            rows.append((idx, title_raw, abs_img))
    except Exception as e:
        print(f"WARN Could not read slides workbook: {workbook_path} ({e})", file=sys.stderr)
    finally:
        try:
            if workbook is not None:
                workbook.close()
        except Exception:
            pass

    rows.sort(key=lambda x: x[0])
    return rows


def build_raw_from_images(
    api_key: str,
    vision_model: str,
    rows: list[tuple[int, str, str]],
    raw_out: str,
    ocr_passes: int,
) -> str:
    raw_text_parts: list[str] = []
    with open(raw_out, "w", encoding="utf-8") as raw_fp:
        for idx, title, abs_img in rows:
            label = f"{idx} - {title}"
            if not os.path.isfile(abs_img):
                block = f"=== SLIDE {idx} ({title}) ===\n[Image file not found: {abs_img}]\n"
                raw_fp.write(block)
                raw_fp.flush()
                raw_text_parts.append(block)
                print(f"WARN missing image: {abs_img}", file=sys.stderr)
                continue
            extracted = vision_extract_with_tiling_retry(
                api_key,
                vision_model,
                abs_img,
                label,
                max_passes=ocr_passes,
            )
            block = f"=== SLIDE {idx} ({title}) ===\n{extracted}\n"
            raw_fp.write(block)
            raw_fp.flush()
            raw_text_parts.append(block)
    return "".join(raw_text_parts)


def build_raw_from_pptx(rows: list[tuple[int, str, str, str, int]], raw_out: str) -> str:
    raw_text_parts: list[str] = []
    with open(raw_out, "w", encoding="utf-8") as raw_fp:
        for idx, title, extracted, _pptx_path, _local_idx in rows:
            block = f"=== SLIDE {idx} ({title}) ===\n{extracted}\n"
            raw_fp.write(block)
            raw_fp.flush()
            raw_text_parts.append(block)
    return "".join(raw_text_parts)


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--root", default=os.getcwd())
    args = p.parse_args()
    root = os.path.normpath(args.root)

    settings_path = os.path.join(root, "Data", "AIPresenter", "Config", "ai_settings.json")
    slides_xlsx_path = os.path.join(root, "Data", "AIPresenter", "Input", "slides.xlsx")
    narration_mode_xlsx_path = os.path.join(root, NARRATION_MODE_XLSX_DEFAULT)
    slide_index_map_xlsx_path = os.path.join(root, SLIDE_INDEX_MAP_XLSX_DEFAULT)
    master_prompt_path = os.path.join(root, "Data", "AIPresenter", "Input", "presentation_master_prompt.txt")
    raw_out = os.path.join(root, "Data", "AIPresenter", "Temp", "extracted_slides_raw.txt")
    gpt_out = os.path.join(root, "Data", "AIPresenter", "Input", "gpt_slide_content.txt")
    os.makedirs(os.path.dirname(raw_out), exist_ok=True)
    os.makedirs(os.path.dirname(gpt_out), exist_ok=True)
    if not os.path.isfile(gpt_out):
        open(gpt_out, "w", encoding="utf-8").close()

    cfg = load_json(settings_path)
    api_key = (cfg.get("openai_api_key") or "").strip()
    vision_model = (cfg.get("openai_vision_model") or "gpt-4o").strip()
    chat_model = (cfg.get("openai_model") or "gpt-4o").strip()
    ocr_passes = parse_ocr_passes(cfg.get("openai_ocr_passes"), default_value=DEFAULT_OCR_PASSES)

    if not api_key or api_key.startswith("YOUR_"):
        print("Missing valid openai_api_key in ai_settings.json", file=sys.stderr)
        sys.exit(1)

    if not os.path.isfile(master_prompt_path):
        print("presentation_master_prompt.txt not found", file=sys.stderr)
        sys.exit(1)

    with open(master_prompt_path, encoding="utf-8") as f:
        master_prompt = f.read().strip()

    configured_pptx_setting = (
        (cfg.get("presentation_pptx_path") or "").strip()
        or (cfg.get("pptx_input_path") or "").strip()
        or (cfg.get("slide_pptx_path") or "").strip()
    )
    xlsx_rows = load_rows_from_xlsx(root, slides_xlsx_path)
    clear_pptx_source_error(root)
    try:
        pptx_paths = resolve_pptx_paths(root, cfg)
    except RuntimeError as e:
        detail = str(e).strip()
        write_pptx_source_error(root, detail)
        print(detail, file=sys.stderr)
        sys.exit(1)

    if configured_pptx_setting and not pptx_paths:
        print(
            "Configured PPTX path not found. Check presentation_pptx_path/pptx_input_path/slide_pptx_path.",
            file=sys.stderr,
        )
        sys.exit(1)

    raw_text = ""
    source_mode = ""
    pptx_slide_lookup: dict[tuple[str, int], int] = {}
    slide_source_meta: dict[int, tuple[str, int]] = {}
    if pptx_paths:
        try:
            all_pptx_rows: list[tuple[int, str, str, str, int]] = []
            next_index = 1
            for pptx_path in pptx_paths:
                deck_rows = load_slides_from_pptx(pptx_path, start_index=next_index)
                all_pptx_rows.extend(deck_rows)
                next_index += len(deck_rows)

            slide_image_map = read_slide_image_map_from_tsv(root)
            if slide_image_map:
                all_pptx_rows = enrich_pptx_rows_with_image_ocr(
                    api_key,
                    vision_model,
                    all_pptx_rows,
                    slide_image_map,
                    ocr_passes=ocr_passes,
                )
        except Exception as e:
            print(f"PPTX parse failed: {e}", file=sys.stderr)
            sys.exit(2)

        raw_text = build_raw_from_pptx(all_pptx_rows, raw_out)
        write_slide_index_mapping_xlsx(root, slide_index_map_xlsx_path, all_pptx_rows)
        pptx_slide_lookup = build_pptx_slide_lookup(root, all_pptx_rows)
        slide_source_meta = build_slide_index_source_meta(all_pptx_rows)
        source_mode = f"pptx:{len(pptx_paths)} file(s)"
        print(
            "INFO pptx order: " + ", ".join(os.path.basename(path) for path in pptx_paths),
            file=sys.stderr,
        )
        print(
            "INFO slide index mapping: " + format_path_for_message(root, slide_index_map_xlsx_path),
            file=sys.stderr,
        )
    elif xlsx_rows:
        try:
            raw_text = build_raw_from_images(
                api_key,
                vision_model,
                xlsx_rows,
                raw_out,
                ocr_passes=ocr_passes,
            )
            source_mode = "xlsx-images"
            if os.path.isfile(slide_index_map_xlsx_path):
                os.remove(slide_index_map_xlsx_path)
        except (urllib.error.HTTPError, urllib.error.URLError, KeyError, OSError) as e:
            print(f"Vision extraction failed: {e}", file=sys.stderr)
            sys.exit(2)
    else:
        print(
            "No slide source found. Provide slides.xlsx or a .pptx in Data/AIPresenter/Input/pptx_file_here "
            "(or set presentation_pptx_path in ai_settings.json).",
            file=sys.stderr,
        )
        sys.exit(1)

    if not raw_text.strip():
        print("No slide content extracted from source.", file=sys.stderr)
        sys.exit(2)

    print(f"INFO OCR multi-pass enabled: {ocr_passes} pass(es)", file=sys.stderr)

    mode_map, mode_warnings = read_slide_narration_mode_xlsx(
        narration_mode_xlsx_path,
        pptx_slide_lookup=pptx_slide_lookup if pptx_slide_lookup else None,
    )
    for warn in mode_warnings:
        print(warn, file=sys.stderr)

    runtime_rules = build_runtime_narration_rules(mode_map)
    mode_user_block = build_narration_mode_user_block(mode_map, slide_source_meta)
    effective_master_prompt = master_prompt + "\n\n" + runtime_rules

    user_block = (
        "RAW DATA (slide order):\n\n"
        + raw_text
        + "\n\n"
        + mode_user_block
        + "\n\nFollow the === SLIDE | index | title === block format as instructed."
    )

    try:
        composed = chat_completion(api_key, chat_model, effective_master_prompt, user_block)
    except (urllib.error.HTTPError, urllib.error.URLError, KeyError) as e:
        print(f"GPT compose failed: {e}", file=sys.stderr)
        sys.exit(3)

    raw_blocks = parse_raw_slide_blocks(raw_text)
    composed = enforce_full_mode_coverage(composed, raw_blocks, mode_map)
    composed = enforce_hyperlink_summary_coverage(composed, raw_blocks)

    with open(gpt_out, "w", encoding="utf-8") as f:
        f.write(composed)

    slides_parsed = parse_slides_blocks(composed)
    if not slides_parsed:
        print(
            "GPT did not return === SLIDE | ... === blocks - UiPath will not find per-slide narration; "
            "check presentation_master_prompt.txt and model output.",
            file=sys.stderr,
        )

    raw_rel = format_path_for_message(root, raw_out)
    gpt_rel = format_path_for_message(root, gpt_out)
    print(f"OK source={source_mode} raw={raw_rel} gpt={gpt_rel} slide_blocks={len(slides_parsed)}")
    sys.exit(0)


if __name__ == "__main__":
    main()
